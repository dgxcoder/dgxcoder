"""Generate the authored part of the ling-docs evaluation corpus: a fictional company's mail (.mbox and
.eml), CSV files, and DOCX files built from collected Wikipedia HTML (CC BY-SA, attributed inside).
Everything here except the DOCX text is written for this set and released under CC0-1.0."""
import csv
import email.utils
import os
import random
import sys
from email.message import EmailMessage

ROOT = sys.argv[1]
os.makedirs(f"{ROOT}/email", exist_ok=True)
os.makedirs(f"{ROOT}/csv", exist_ok=True)
os.makedirs(f"{ROOT}/docx", exist_ok=True)

PEOPLE = {
    "maria": "Maria Okafor <maria.okafor@kestrel-robotics.example>",
    "jonas": "Jonas Lindqvist <jonas.lindqvist@kestrel-robotics.example>",
    "priya": "Priya Raman <priya.raman@kestrel-robotics.example>",
    "tomas": "Tomas Herrera <tomas.herrera@kestrel-robotics.example>",
    "acme": "Daniel Brooks <d.brooks@acme-actuators.example>",
    "fjord": "Ingrid Solberg <ingrid@fjordfreight.example>",
    "audit": "Chen Wei <chen.wei@harbourline-audit.example>",
}

# (from, to, subject, day, body). Facts here are what the questions ask about.
THREADS = [
    ("acme", "maria", "Warranty on the RX-40 servo batch", 3,
     "Hi Maria,\n\nFollowing our call, Acme Actuators agrees to extend the warranty on the RX-40 servo batch "
     "(order 7731) from 12 to 24 months at no extra cost, because of the bearing noise you reported. The "
     "extension applies to units shipped between March and June.\n\nRegards,\nDaniel"),
    ("maria", "acme", "Re: Warranty on the RX-40 servo batch", 3,
     "Daniel,\n\nThank you. Please send the amended warranty certificate to procurement before the end of the "
     "month so finance can close the claim.\n\nMaria"),
    ("fjord", "jonas", "Delay: container KRB-2291", 5,
     "Hello Jonas,\n\nThe container with the gripper assemblies (KRB-2291) is held in Rotterdam because of a "
     "customs documentation error: the HS code on the commercial invoice was 8479 instead of 8428. We expect "
     "release within four working days once a corrected invoice is filed.\n\nIngrid, Fjord Freight"),
    ("jonas", "fjord", "Re: Delay: container KRB-2291", 6,
     "Ingrid, corrected invoice attached. Please confirm the new arrival date for the Gothenburg warehouse.\nJonas"),
    ("fjord", "jonas", "Re: Re: Delay: container KRB-2291", 9,
     "Jonas, customs released KRB-2291 this morning. New arrival in Gothenburg: the 19th, morning slot.\nIngrid"),
    ("priya", "maria", "Q3 hiring plan", 10,
     "Maria,\n\nFor Q3 I am proposing three hires: two firmware engineers for the motion-control team and one "
     "test technician for the Gothenburg lab. Total budget impact is 412,000 EUR annualised. I would like to "
     "open the firmware roles before the summer break.\n\nPriya"),
    ("maria", "priya", "Re: Q3 hiring plan", 11,
     "Priya, approved for the two firmware roles. The technician role waits until the lab lease is signed.\nMaria"),
    ("tomas", "maria", "Security incident report: build server", 12,
     "Maria,\n\nSummary of yesterday's incident: a leaked deploy token was used from an unknown IP to read the "
     "artifact store of the build server. No source code was modified. We revoked the token within 40 minutes, "
     "rotated all CI secrets, and moved artifact access behind the VPN. Root cause: the token was committed to a "
     "public test repository by mistake.\n\nTomas"),
    ("maria", "tomas", "Re: Security incident report: build server", 12,
     "Tomas, thank you. Please schedule the post-mortem for Thursday and include the contractor who owned the "
     "test repository.\nMaria"),
    ("audit", "maria", "Audit fieldwork dates", 14,
     "Dear Maria,\n\nHarbourline Audit will carry out the year-end fieldwork from 2 to 6 February at your Malmo "
     "office. Please prepare the fixed-asset register and the inventory count sheets for the Gothenburg "
     "warehouse in advance.\n\nKind regards,\nChen Wei"),
    ("jonas", "maria", "Office move to Hamngatan", 15,
     "Maria,\n\nThe landlord confirmed the Hamngatan 14 space from 1 October. Rent is 18,500 SEK per month, "
     "three-year term, with the first two months free in exchange for us paying the fit-out of the meeting "
     "rooms.\n\nJonas"),
    ("maria", "jonas", "Re: Office move to Hamngatan", 16,
     "Jonas, fine by me. Book the movers for the last weekend of September.\nMaria"),
    ("acme", "maria", "Price increase notice", 17,
     "Maria,\n\nPlease note that from 1 January the list price of the RX-40 servo rises by 6 percent, and the "
     "RX-60 by 4 percent, due to higher magnet costs. Orders placed before 15 December keep current pricing.\n\n"
     "Daniel"),
    ("priya", "tomas", "Lab power upgrade", 18,
     "Tomas,\n\nThe Gothenburg lab needs a dedicated 32 A three-phase circuit for the new environmental "
     "chamber. The electrician can do it on a Saturday; estimated cost 23,000 SEK.\n\nPriya"),
    ("tomas", "priya", "Re: Lab power upgrade", 19,
     "Priya, go ahead, charge it to the facilities budget, not R&D.\nTomas"),
    ("maria", "jonas", "Board meeting agenda", 20,
     "Jonas,\n\nAgenda for the board meeting: Q3 results, the Acme warranty settlement, the security incident and "
     "the decision on opening a sales office in Oslo. Please prepare the cash-flow forecast for the Oslo case "
     "with a break-even no later than month eighteen.\n\nMaria"),
    ("jonas", "maria", "Re: Board meeting agenda", 21,
     "Maria, forecast done: the Oslo office breaks even in month fourteen under the base case.\nJonas"),
    ("fjord", "jonas", "New rate card", 22,
     "Jonas,\n\nOur new rate card is attached. Full container Rotterdam to Gothenburg is 1,840 EUR; express "
     "pallet is 295 EUR. Rates are valid until 31 March.\n\nIngrid"),
    ("priya", "maria", "Firmware release 3.2 slipped", 23,
     "Maria,\n\nFirmware 3.2 for the motion controller slips by two weeks. The torque-limit regression on the "
     "RX-60 axis needs a fix in the PID tuning tables before we can sign off the safety tests.\n\nPriya"),
    ("tomas", "maria", "Laptop refresh", 24,
     "Maria,\n\nThe laptop refresh covers 27 machines older than four years. We pick the 14-inch model with 32 GB "
     "of memory; delivery in two batches, half in November and half in January.\n\nTomas"),
    ("acme", "maria", "Settlement proposal", 25,
     "Maria,\n\nTo settle the RX-40 bearing claim, Acme proposes a credit note of 38,000 EUR against future "
     "orders, in addition to the extended warranty already agreed.\n\nDaniel"),
    ("maria", "acme", "Re: Settlement proposal", 26,
     "Daniel, we accept the credit note on condition that it can also be used for spare parts.\nMaria"),
    ("audit", "maria", "Findings: inventory count", 27,
     "Maria,\n\nThe inventory count found a difference of 112 gripper fingers between the warehouse system and "
     "the physical count, valued at 6,720 EUR. We recommend a monthly cycle count for high-value small parts.\n\n"
     "Chen Wei"),
    ("priya", "jonas", "Customer visit from Volvo Cars", 28,
     "Jonas,\n\nThe Volvo Cars delegation visits the Gothenburg lab on the 12th. They want to see the welding cell "
     "demo and discuss a pilot of 40 robots for their Torslanda plant.\n\nPriya"),
    ("jonas", "priya", "Re: Customer visit from Volvo Cars", 28,
     "Priya, I will join. Please make sure the demo uses the 3.1 firmware, not the 3.2 beta.\nJonas"),
    ("tomas", "maria", "Backup policy change", 29,
     "Maria,\n\nFrom next month, backups of the engineering file server run nightly to the Malmo office and weekly "
     "to an offline disk kept in the safe. Retention is 90 days for nightly and one year for weekly copies.\n\n"
     "Tomas"),
    ("maria", "priya", "Patent application", 30,
     "Priya,\n\nThe patent attorney filed our application for the adaptive gripper force sensing method last "
     "week. We have twelve months to decide on international filing under the PCT.\n\nMaria"),
    ("fjord", "jonas", "Peak season surcharge", 31,
     "Jonas,\n\nA peak season surcharge of 120 EUR per container applies from 15 November to 31 December.\n\nIngrid"),
    ("priya", "maria", "Conference talk accepted", 32,
     "Maria,\n\nOur talk on vibration damping for lightweight arms was accepted at the Automatica fair in Munich. "
     "I will present on the second day.\n\nPriya"),
    ("jonas", "maria", "Insurance renewal", 33,
     "Maria,\n\nThe product liability insurance renews on 1 December. The premium goes up 9 percent to 61,000 SEK "
     "because of the larger robot fleet in the field.\n\nJonas"),
]
EML_EXTRA = [
    ("acme", "jonas", "RX-60 lead times", 40,
     "Jonas,\n\nCurrent lead time for the RX-60 is eleven weeks; the RX-40 is in stock.\n\nDaniel"),
    ("tomas", "priya", "Test rig network", 41,
     "Priya,\n\nThe test rigs move to their own VLAN 220, isolated from the office network. Remote access goes "
     "through the jump host only.\n\nTomas"),
    ("audit", "maria", "Draft management letter", 42,
     "Maria,\n\nOur draft management letter has two points: segregation of duties in purchasing, and the missing "
     "review of user access to the ERP system.\n\nChen Wei"),
    ("priya", "tomas", "Oscilloscope purchase", 43,
     "Tomas,\n\nWe need a four-channel 500 MHz oscilloscope for the motor-drive debugging; quotes are 9,800 and "
     "11,200 EUR.\n\nPriya"),
    ("maria", "jonas", "Holiday closure", 44,
     "Jonas,\n\nThe offices close from 23 December to 2 January; the Gothenburg warehouse keeps a skeleton crew.\n\n"
     "Maria"),
    ("fjord", "jonas", "Lost pallet claim", 45,
     "Jonas,\n\nWe accept liability for the missing pallet on shipment FF-88412 and will refund 2,350 EUR.\n\nIngrid"),
    ("jonas", "maria", "Bank covenant", 46,
     "Maria,\n\nThe bank requires the equity ratio to stay above 30 percent at each quarter end under the new "
     "credit facility.\n\nJonas"),
    ("priya", "maria", "Intern project", 47,
     "Maria,\n\nThe summer intern built a dashboard for gripper wear, predicting replacement from cycle counts.\n\n"
     "Priya"),
    ("tomas", "maria", "Password manager rollout", 48,
     "Maria,\n\nThe company password manager is rolled out to all staff; shared vault for the production "
     "credentials is limited to four people.\n\nTomas"),
    ("acme", "maria", "Factory audit invitation", 49,
     "Maria,\n\nYou are welcome to audit our servo production line in Brno in the first week of March.\n\nDaniel"),
]


def message(i, frm, to, subject, day, body):
    m = EmailMessage()
    m["From"] = PEOPLE[frm]
    m["To"] = PEOPLE[to]
    m["Subject"] = subject
    m["Date"] = email.utils.format_datetime(email.utils.parsedate_to_datetime(
        f"Mon, {1 + day % 28:02d} Sep 2025 09:{i % 60:02d}:00 +0200"))
    m["Message-ID"] = f"<kr-{i:03d}@kestrel-robotics.example>"
    m.set_content(body)
    return m


with open(f"{ROOT}/email/kestrel-2025.mbox", "w") as box:
    for i, t in enumerate(THREADS):
        m = message(i, *t)
        box.write(f"From {email.utils.parseaddr(m['From'])[1]} Mon Sep  1 09:00:00 2025\n")
        box.write(m.as_string().replace("\nFrom ", "\n>From ") + "\n")
for j, t in enumerate(EML_EXTRA):
    open(f"{ROOT}/email/kestrel-{j + 100:03d}.eml", "w").write(message(j + 100, *t).as_string())

rng = random.Random(20261007)
with open(f"{ROOT}/csv/warehouse-inventory.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["sku", "description", "location", "quantity", "unit_cost_eur"])
    parts = ["gripper finger", "servo bracket", "cable harness", "encoder disk", "bearing 6202", "torque sensor",
             "end stop", "belt GT2", "pulley 20T", "motor mount"]
    for k in range(200):
        p = parts[k % len(parts)]
        w.writerow([f"KR-{1000 + k}", f"{p} rev {k // 10}", f"GBG-{rng.randint(1, 40):02d}-{rng.randint(1, 9)}",
                    rng.randint(0, 900), round(rng.uniform(1, 400), 2)])
    w.writerow(["KR-1999", "safety relay module", "GBG-07-3", 3, 412.5])
with open(f"{ROOT}/csv/expenses-2025.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["date", "department", "category", "vendor", "amount_sek"])
    cats = [("R&D", "components"), ("R&D", "software licences"), ("Facilities", "rent"), ("Facilities", "cleaning"),
            ("Sales", "travel"), ("Sales", "trade fairs"), ("IT", "laptops"), ("IT", "cloud")]
    for k in range(240):
        d, c = cats[k % len(cats)]
        w.writerow([f"2025-{1 + k % 12:02d}-{1 + k % 28:02d}", d, c, f"vendor-{rng.randint(1, 60)}",
                    round(rng.uniform(500, 90000), 2)])
    w.writerow(["2025-11-14", "Facilities", "electrical work", "Elteknik Goteborg AB", 23000.0])
with open(f"{ROOT}/csv/robot-fleet.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["serial", "model", "customer", "site", "firmware", "commissioned"])
    for k in range(150):
        w.writerow([f"KR{24000 + k}", rng.choice(["K-arm 5", "K-arm 7", "K-weld 2"]),
                    rng.choice(["Scania", "SKF", "Ericsson", "Husqvarna", "ABB Service"]),
                    rng.choice(["Sodertalje", "Gothenburg", "Kista", "Huskvarna", "Vasteras"]),
                    rng.choice(["3.0", "3.1"]), f"202{rng.randint(2, 5)}-{rng.randint(1, 12):02d}"])
    w.writerow(["KR24999", "K-weld 2", "Volvo Cars", "Torslanda", "3.1", "2025-09"])
with open(f"{ROOT}/csv/sensor-calibration.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["sensor_id", "type", "offset", "gain", "calibrated_on", "technician"])
    for k in range(120):
        w.writerow([f"TS-{k:04d}", rng.choice(["torque", "force", "temperature"]), round(rng.uniform(-2, 2), 4),
                    round(rng.uniform(0.95, 1.05), 5), f"2025-0{rng.randint(1, 9)}-{rng.randint(10, 28)}",
                    rng.choice(["A. Berg", "L. Nilsson", "K. Haddad"])])
with open(f"{ROOT}/csv/world-capitals.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["country", "capital", "continent"])
    for row in [("Iceland", "Reykjavik", "Europe"), ("Norway", "Oslo", "Europe"), ("Sweden", "Stockholm", "Europe"),
                ("Finland", "Helsinki", "Europe"), ("Denmark", "Copenhagen", "Europe"),
                ("Estonia", "Tallinn", "Europe"), ("Latvia", "Riga", "Europe"), ("Lithuania", "Vilnius", "Europe"),
                ("Portugal", "Lisbon", "Europe"), ("Chile", "Santiago", "South America"),
                ("Peru", "Lima", "South America"), ("Kenya", "Nairobi", "Africa"), ("Ghana", "Accra", "Africa"),
                ("Japan", "Tokyo", "Asia"), ("Mongolia", "Ulaanbaatar", "Asia"), ("Bhutan", "Thimphu", "Asia"),
                ("New Zealand", "Wellington", "Oceania"), ("Canada", "Ottawa", "North America")]:
        w.writerow(row)

# DOCX from the collected Wikipedia HTML (section headings + paragraphs).
from bs4 import BeautifulSoup  # noqa: E402
from docx import Document  # noqa: E402

src = f"{ROOT}/_docx_src"
for name in sorted(os.listdir(src)) if os.path.isdir(src) else []:
    soup = BeautifulSoup(open(f"{src}/{name}", encoding="utf-8").read(), "html.parser")
    for tag in soup.select("table, style, sup.reference, .mw-ref, figure, .hatnote, .navbox"):
        tag.decompose()
    doc = Document()
    title = name[:-5].replace("_", " ")
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Text adapted from the Wikipedia article \"{title}\" by Wikipedia contributors, "
                      "CC BY-SA 4.0. Generated for the ling-docs evaluation set.")
    for el in soup.find_all(["h2", "h3", "p"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name == "h2":
            doc.add_heading(text, level=1)
        elif el.name == "h3":
            doc.add_heading(text, level=2)
        else:
            doc.add_paragraph(text)
    doc.core_properties.author = "Wikipedia contributors"
    doc.save(f"{ROOT}/docx/{name[:-5]}.docx")
print("generated")
